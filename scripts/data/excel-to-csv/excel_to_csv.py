"""Export the sheets of an Excel workbook (.xlsx) to CSV files."""
import argparse
import csv
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def safe_name(text):
    return re.sub(r"[^\w.-]+", "_", text).strip("_") or "sheet"


def cell_text(value):
    """Plain text for a cell: dates as ISO, whole floats without '.0', None as ''."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():  # Excel stores plain dates as midnight
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def sheet_rows(sheet):
    """Rows as lists of text, with fully empty trailing rows dropped."""
    rows = [[cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def main(argv=None):
    parser = argparse.ArgumentParser(description="Export Excel sheets to CSV.")
    parser.add_argument("file", type=Path, help=".xlsx workbook")
    parser.add_argument("--sheet", help="only this sheet (default: all sheets)")
    parser.add_argument("-o", "--outdir", type=Path, help="folder for the CSV files")
    parser.add_argument("--apply", action="store_true",
                        help="really write the CSV files (default is a dry run)")
    args = parser.parse_args(argv)
    if not args.file.is_file():
        parser.error(f"not a file: {args.file}")

    workbook = load_workbook(args.file, read_only=True, data_only=True)
    names = [args.sheet] if args.sheet else workbook.sheetnames
    unknown = [n for n in names if n not in workbook.sheetnames]
    if unknown:
        parser.error(f"no such sheet: {unknown[0]} (available: {', '.join(workbook.sheetnames)})")

    outdir = args.outdir or args.file.parent
    for name in names:
        rows = sheet_rows(workbook[name])
        target = outdir / f"{args.file.stem}_{safe_name(name)}.csv"
        print(f"{name}: {len(rows)} row(s) -> {target}")
        if args.apply:
            outdir.mkdir(parents=True, exist_ok=True)
            with target.open("w", newline="", encoding="utf-8") as handle:
                csv.writer(handle).writerows(rows)
    verb = "Wrote" if args.apply else "Would write"
    hint = "" if args.apply else " Use --apply to do it."
    print(f"{verb} {len(names)} file(s).{hint}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
